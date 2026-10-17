"""Blind label sheets: an annotator's queue as Excel, and the filled-in sheet read back
as their judgements, with every problem reported by row."""

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app import cli
from app.core.config import get_settings
from app.models.event import EventEvidence, EvidenceJudgement, EvidenceRelation
from app.models.user import Role
from app.services import annotation, label_sheet
from tests.conftest import TEST_DATABASE_URL
from tests.test_assessment import _event, _feed, _report, _signal
from tests.test_assessment import people as people  # noqa: F401 (fixture)

EMAIL = "annotator.b@carcux.org"


@pytest.fixture
def test_db_cli(monkeypatch):
    """The commands open their own connection: point them at the test database."""
    monkeypatch.setenv("CARCUX_DATABASE_URL", TEST_DATABASE_URL)
    get_settings.cache_clear()
    yield
    get_settings.cache_clear()


def _setup(db, people, make_user):
    a, (w1, w2, _) = people["analyst"], people["workers"]
    e = _event(db, a)
    _report(db, e, w1, a, relation=EvidenceRelation.CONTRADICTS, text="Mirpur 10 e pani, gujob")
    _report(db, e, w2, a, text="Knee-deep water on Mirpur 10 road", minutes=5)
    _signal(db, e, _feed(db), a, minutes=10)
    db.flush()
    for link in db.scalars(select(EventEvidence)):
        db.add(EvidenceJudgement(evidence_id=link.id, annotator_id=a.id, relation=link.relation))
    db.commit()
    return make_user(EMAIL, Role.ANNOTATOR, full_name="Annotator B")


def _fill(path, answers):
    wb = load_workbook(path)
    ws = wb["Items"]
    for row, values in answers.items():
        for offset, value in enumerate(values):
            ws.cell(row=row, column=label_sheet.FIRST_ANSWER + offset, value=value)
    wb.save(path)


def _rows_by_text(path) -> dict[str, int]:
    """The sheet row of each test item, found by a word in what it says."""
    ws = load_workbook(path)["Items"]
    out = {}
    for r in range(2, ws.max_row + 1):
        says = ws.cell(row=r, column=6).value
        for key in ("gujob", "Knee-deep", "Mirpur flooded"):
            if key in says:
                out[key] = r
    return out


def test_the_sheet_is_the_blind_queue(db, people, make_user, tmp_path):
    user = _setup(db, people, make_user)
    path = tmp_path / "sheet.xlsx"
    assert label_sheet.export(db, user, path) == 3
    wb = load_workbook(path)
    assert wb.sheetnames == ["How to label", "Items"]
    ws = wb["Items"]
    cells = [str(c.value) for row in ws.iter_rows() for c in row if c.value is not None]
    text = " ".join(cells).lower()
    # Nobody's label and nothing of CARCUX's reading; the answer columns are empty.
    for leak in ("contradicts", "partially_supports", "stance", "assessment", "test user"):
        assert leak not in text.replace("contradicts:", ""), leak
    assert all(ws.cell(row=r, column=label_sheet.FIRST_ANSWER).value is None for r in (2, 3, 4))
    id_letter = ws.cell(row=1, column=label_sheet.ID_COL).column_letter
    assert ws.column_dimensions[id_letter].hidden
    assert "Mirpur 10 e pani, gujob" in text or "mirpur 10 e pani, gujob" in text


def test_a_filled_sheet_becomes_the_annotators_judgements(db, people, make_user, tmp_path):
    user = _setup(db, people, make_user)
    path = tmp_path / "sheet.xlsx"
    label_sheet.export(db, user, path)
    row = _rows_by_text(path)
    _fill(
        path,
        {
            row["gujob"]: ["Contradicts", None, "No", "Certain", "calls it gujob"],
            row["Knee-deep"]: ["Partly supports", "Deaths, place", "yes", "Unsure"],
            row["Mirpur flooded"]: ["related only"],
        },
    )
    sheet = label_sheet.read(db, path)
    assert sheet.errors == [] and len(sheet.rows) == 3
    assert label_sheet.save(db, user, sheet) == 3
    mine = db.scalars(select(EvidenceJudgement).where(EvidenceJudgement.annotator_id == user.id))
    got = sorted((j.relation.value, tuple(j.conflicts), j.stale, j.confidence) for j in mine)
    assert got == [
        ("contradicts", (), False, 3),
        ("partially_supports", ("casualty_count", "location"), True, 1),
        ("related", (), False, 2),
    ]
    stats = annotation.agreement(db)
    assert stats.items_with_two == 3 and stats.percent == 1 / 3


def test_problems_are_reported_by_row_and_nothing_is_saved(
    db, people, make_user, tmp_path, test_db_cli
):
    user = _setup(db, people, make_user)
    path = tmp_path / "sheet.xlsx"
    label_sheet.export(db, user, path)
    _fill(path, {2: ["Maybe"], 3: ["Partly supports"], 4: ["Partly supports", "Weather"]})
    sheet = label_sheet.read(db, path)
    assert len(sheet.errors) == 3
    assert "row 2" in sheet.errors[0] and "unknown label 'Maybe'" in sheet.errors[0]
    assert "partly supports needs a conflict" in sheet.errors[1]
    assert "weather" in sheet.errors[2]

    # Through the command: checked, refused, nothing recorded.
    assert cli.main(["import-label-sheet", EMAIL, str(path), "--save"]) == 1
    count = db.scalars(select(EvidenceJudgement).where(EvidenceJudgement.annotator_id == user.id))
    assert list(count) == []


def test_blank_rows_wait_and_a_changed_id_is_caught(db, people, make_user, tmp_path):
    user = _setup(db, people, make_user)
    path = tmp_path / "sheet.xlsx"
    label_sheet.export(db, user, path)
    wb = load_workbook(path)
    wb["Items"].cell(row=2, column=label_sheet.ID_COL, value="not-an-id")
    wb["Items"].cell(row=3, column=label_sheet.FIRST_ANSWER, value="Supports")
    wb.save(path)
    sheet = label_sheet.read(db, path)
    assert sheet.blank == 1 and len(sheet.rows) == 1
    assert "hidden id was changed" in sheet.errors[0]


def test_the_commands(db, people, make_user, tmp_path, capsys, test_db_cli):
    _setup(db, people, make_user)
    path = tmp_path / "sheet.xlsx"
    assert cli.main(["export-label-sheet", EMAIL, str(path)]) == 0
    assert "Wrote 3 items for Annotator B" in capsys.readouterr().out
    _fill(path, {2: ["Supports"], 3: ["Supports"], 4: ["Supports"]})
    assert cli.main(["import-label-sheet", EMAIL, str(path)]) == 0
    assert "Checked only" in capsys.readouterr().out
    assert cli.main(["import-label-sheet", EMAIL, str(path), "--save"]) == 0
    assert "Saved 3 judgements" in capsys.readouterr().out
    assert cli.main(["export-label-sheet", "nobody@carcux.org", str(path)]) == 1
