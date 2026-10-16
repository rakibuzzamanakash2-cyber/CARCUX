"""Blind label sheets: second opinions given away from CARCUX (guideline section 7).

An annotator who cannot use the console gets their blind queue as an Excel workbook:
the event as it is judged against (what, where, when, ground truth) and the item, with
a dropdown for the label. Like the queue, it carries no one's label and nothing of
CARCUX's reading. The filled-in sheet is read back as that annotator's judgements,
exactly as if they had used the queue.
"""

import re
import uuid
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy.orm import Session

from app.models.event import ClaimAttribute, EventEvidence, EvidenceRelation
from app.models.user import User
from app.services import annotation
from app.services.details import headline_and_text

LABELS = {
    "Supports": EvidenceRelation.SUPPORTS,
    "Partly supports": EvidenceRelation.PARTIALLY_SUPPORTS,
    "Contradicts": EvidenceRelation.CONTRADICTS,
    "Related only": EvidenceRelation.RELATED,
    "Not this event": EvidenceRelation.UNRELATED,
}
# The words in the console's "What does it get wrong?" list.
WRONG = {
    "kind of event": ClaimAttribute.EVENT_TYPE,
    "place": ClaimAttribute.LOCATION,
    "when it started": ClaimAttribute.START_TIME,
    "when it ended": ClaimAttribute.END_TIME,
    "still going on": ClaimAttribute.STATUS,
    "whether it is still going on": ClaimAttribute.STATUS,
    "how bad": ClaimAttribute.MAGNITUDE,
    "number affected": ClaimAttribute.AFFECTED_COUNT,
    "deaths": ClaimAttribute.CASUALTY_COUNT,
    "cause": ClaimAttribute.CAUSE,
}
WRONG_HELP = (
    "Deaths, Number affected, Place, When it started, When it ended, Still going on, "
    "How bad, Cause, Kind of event"
)
SURE = {"Unsure": 1, "Fairly sure": 2, "Certain": 3}

COLUMNS = [
    ("No.", 6),
    ("Event", 34),
    ("Event: where and when", 26),
    ("Ground truth", 46),
    ("Item: source and date", 24),
    ("Item: what it says", 60),
    ("Link to the item", 30),
    ("Your label", 18),
    ("If Partly supports: what it gets wrong", 30),
    ("Out of date when published?", 14),
    ("How sure", 14),
    ("Note (optional)", 30),
    ("id (do not change)", 38),
]
ID_COL = len(COLUMNS)  # last column, hidden
FIRST_ANSWER = 8  # "Your label"

GUIDE = [
    "How to label (CARCUX-BD pilot, guideline v0.2)",
    "",
    "Each row is one news item or post, and the event it may be about. Say what the item "
    "says about that event. Another person has labelled the same items; your labels must "
    "be your own, so do not discuss items with anyone until you have finished.",
    "",
    "Label what the item says, not whether it is true. A post claiming a collapse that "
    "never happened still supports that rumoured event.",
    "You may open an item's link to read it in full. Do not search the web for the event.",
    "",
    "For each row, read the event and its ground truth, then the item. Ask in this order "
    "and stop at the first yes:",
    "1. Is it about a different event (wrong place, day or type)?  →  Not this event",
    "2. Does it make no checkable claim (prayer, opinion, complaint, aftermath talk)?  →  "
    "Related only",
    "3. Does it deny the event, or call it false or a rumour?  →  Contradicts",
    "4. Does it say it happened, but get something wrong (a count or status that was wrong "
    "when it was published, the wrong place, depth or cause)?  →  Partly supports, and "
    "write what is wrong",
    "5. Otherwise  →  Supports",
    "",
    "Counts agree if they are within ±20% of the official toll known when the item was "
    "published (the ground truth says how the toll rose). A toll that was right that day "
    "is Supports, even if it rose later.",
    "What it gets wrong: one or more of " + WRONG_HELP + ", separated by commas.",
    "Out of date when published: Yes if the item describes an earlier moment that had "
    "already passed when it went up (an old photo, 'still underwater' after the water "
    "had gone).",
    "How sure: Unsure is fine and useful. Note: a few words on what decided a hard one.",
    "",
    "Fill in the four columns 'Your label' to 'How sure' (and a note if you like). Do not "
    "change the other columns or the order of the rows. Send the file back when every "
    "row has a label.",
]


def _when(dt) -> str:
    return dt.strftime("%d %b %Y, %H:%M") if dt else ""


def _row(n: int, link: EventEvidence) -> list:
    e = link.event
    truth = (
        "not established yet"
        if e.occurred is None
        else ("it happened" if e.occurred else "it did not happen")
    )
    if e.ground_truth_note:
        truth += ". " + e.ground_truth_note
    where_when = f"{e.place_name or ''}\n{_when(e.started_at)}" + (
        f" to {_when(e.ended_at)}" if e.ended_at else ", ongoing"
    )
    if link.signal:
        s = link.signal
        source = f"{s.source.name}\n{_when(s.published_at)}"
        says = headline_and_text(s.title, s.text)
        url = s.url or ""
    else:
        r = link.field_report
        source = f"Field report\n{_when(r.observed_at)}"
        says, url = r.text, ""
    title = f"{e.title}\n({e.event_type.replace('_', ' ')})"
    return [n, title, where_when, truth, source, says, url, None, None, None, None, None,
            str(link.id)]  # fmt: skip


def export(db: Session, annotator: User, path: Path) -> int:
    """Write `annotator`'s blind queue to `path`; returns the number of rows."""
    links, _ = annotation.queue(db, annotator, limit=10_000, offset=0)
    wb = Workbook()
    guide = wb.active
    guide.title = "How to label"
    guide.column_dimensions["A"].width = 120
    for i, line in enumerate(GUIDE, start=1):
        cell = guide.cell(row=i, column=1, value=line)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        if i == 1:
            cell.font = Font(bold=True, size=14)

    ws = wb.create_sheet("Items")
    head = PatternFill("solid", fgColor="1F5130")
    answer = PatternFill("solid", fgColor="FFF4D6")
    for c, (name, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=c, value=name)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = head
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions[cell.column_letter].width = width
    for n, link in enumerate(links, start=1):
        for c, value in enumerate(_row(n, link), start=1):
            cell = ws.cell(row=n + 1, column=c, value=value)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if FIRST_ANSWER <= c < ID_COL:
                cell.fill = answer
    ws.freeze_panes = "B2"
    ws.column_dimensions[ws.cell(row=1, column=ID_COL).column_letter].hidden = True
    last = len(links) + 1
    for col, options in (
        (FIRST_ANSWER, list(LABELS)),
        (FIRST_ANSWER + 2, ["Yes", "No"]),
        (FIRST_ANSWER + 3, list(SURE)),
    ):
        dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
        letter = ws.cell(row=1, column=col).column_letter
        dv.add(f"{letter}2:{letter}{max(last, 2)}")
        ws.add_data_validation(dv)
    wb.active = ws
    wb.save(path)
    return len(links)


@dataclass
class Read:
    rows: list[dict] = field(default_factory=list)  # ready to save
    errors: list[str] = field(default_factory=list)
    blank: int = 0


def _clean(value) -> str:
    return re.sub(r"\s+", " ", str(value)).strip() if value is not None else ""


def read(db: Session, path: Path) -> Read:
    """Check a filled-in sheet; every problem is reported with its row number."""
    out = Read()
    ws = load_workbook(path, read_only=True, data_only=True)["Items"]
    for i, values in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        values = list(values) + [None] * (ID_COL - len(values))
        if not any(values):
            continue
        label, wrong, stale, sure, note, raw_id = values[FIRST_ANSWER - 1 : ID_COL]
        where = f"row {i} (No. {values[0]})"
        try:
            evidence_id = uuid.UUID(_clean(raw_id))
        except ValueError:
            out.errors.append(f"{where}: the hidden id was changed or removed")
            continue
        if db.get(EventEvidence, evidence_id) is None:
            out.errors.append(f"{where}: this link is no longer in CARCUX")
            continue
        label = _clean(label)
        if not label:
            out.blank += 1
            continue
        relation = next((v for k, v in LABELS.items() if k.lower() == label.lower()), None)
        if relation is None:
            out.errors.append(f"{where}: unknown label '{label}'")
            continue
        conflicts, unknown = [], []
        for word in filter(None, (_clean(w).lower() for w in re.split(r"[,;/\n]", _clean(wrong)))):
            if word in WRONG:
                conflicts.append(WRONG[word].value)
            else:
                unknown.append(word)
        if unknown:
            out.errors.append(f"{where}: unknown in 'what it gets wrong': {', '.join(unknown)}")
            continue
        conflicts = list(dict.fromkeys(conflicts))
        if relation != EvidenceRelation.PARTIALLY_SUPPORTS:
            conflicts = []  # only partly supports names what is wrong
        try:
            annotation.check_labels(relation, conflicts)
        except annotation.InvalidJudgementError as exc:
            out.errors.append(f"{where}: {exc}")
            continue
        sure_value = _clean(sure)
        if sure_value and sure_value not in SURE:
            out.errors.append(f"{where}: 'How sure' must be Unsure, Fairly sure or Certain")
            continue
        out.rows.append(
            {
                "evidence_id": evidence_id,
                "relation": relation,
                "conflicts": conflicts,
                "stale": _clean(stale).lower() in {"yes", "y", "true"},
                "confidence": SURE.get(sure_value, 2),
                "note": _clean(note) or None,
            }
        )
    return out


def save(db: Session, annotator: User, sheet: Read) -> int:
    """Record each row as `annotator`'s judgement (replacing an earlier one of theirs)."""
    for row in sheet.rows:
        annotation.judge(
            db,
            actor=annotator,
            item=db.get(EventEvidence, row["evidence_id"]),
            relation=row["relation"],
            conflicts=row["conflicts"],
            stale=row["stale"],
            confidence=row["confidence"],
            note=row["note"],
            ip_address=None,
        )
    return len(sheet.rows)
