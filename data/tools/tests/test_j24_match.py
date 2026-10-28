"""The matcher on SYNTHETIC claims and articles: Bangla stems, the date window, ranking,
the snippet, and the candidate sheet (which must not show verdicts)."""

import json
from datetime import date

from openpyxl import load_workbook

from carcux_data.j24 import match as m


def test_light_bangla_stemming_joins_case_endings():
    assert m.stem("শিক্ষার্থীদের") == m.stem("শিক্ষার্থী") == "শিক্ষার্থী"
    assert m.stem("police") == "police"
    assert "দাবি" not in m.tokens("দাবি করা হয়েছে")


def write(tmp_path, claims, articles):
    (tmp_path / "news").mkdir()
    (tmp_path / "news" / "texts.jsonl").write_text(
        "".join(json.dumps(a, ensure_ascii=False) + "\n" for a in articles), encoding="utf-8"
    )
    path = tmp_path / "claims.jsonl"
    path.write_text(
        "".join(json.dumps(c, ensure_ascii=False) + "\n" for c in claims), encoding="utf-8"
    )
    return path


def claim(cid, text, day, group=None, **kw):
    return {
        "claim_id": cid,
        "claim_group": group or cid,
        "in_scope": "yes",
        "claim": text,
        "event_date": day,
        "published_at": day,
        "source": "fact_check",
        "fact_checker": "X",
        "title": kw.get("title", ""),
        "verdict": "false",
        "evidence_text": kw.get("evidence", ""),
    }


def art(url, title, text, day):
    return {
        "url": url,
        "outlet": "paper",
        "title": title,
        "text": text,
        "published_at": day,
        "captured": None,
    }


def test_ranks_the_matching_article_first_and_keeps_to_the_window(tmp_path):
    claims = [
        claim(
            "c1",
            "seven people were killed at Kotbari bridge",
            "2024-07-18",
            title="কোটবাড়ি সেতুতে সাত জন নিহত",
        )
    ]
    arts = [
        art(
            "u1",
            "কোটবাড়িতে সংঘর্ষ",
            "SYNTHETIC কোটবাড়ি সেতুতে সংঘর্ষে সাত জন নিহত হয়েছেন " * 5,
            "2024-07-19",
        ),
        art("u2", "বাজারদর", "SYNTHETIC চালের দাম বেড়েছে " * 10, "2024-07-19"),
        art("u3", "কোটবাড়ি", "SYNTHETIC কোটবাড়ি সেতুতে সাত জন নিহত " * 5, "2024-05-01"),
    ]
    path = write(tmp_path, claims, arts)
    res = m.match(tmp_path, path, top=5, log=lambda s: None)
    assert [r["url"] for r in res] == ["u1"]  # u2 shares no words, u3 is out of the window
    assert "কোটবাড়ি" in res[0]["snippet"]


def test_linked_rows_form_one_query(tmp_path):
    claims = [
        claim("c1", "SYNTHETIC claim about the airport", "2024-07-16"),
        claim("c2", "", "2024-07-16", group="c1", title="বিমানবন্দর বন্ধ"),
    ]
    groups = m.claim_groups(write(tmp_path, claims, []))
    assert len(groups) == 1 and "বিমানবন্দর" in groups[0]["query"]
    assert groups[0]["ref_date"] == date(2024, 7, 16)


def test_candidate_sheet_has_no_verdict(tmp_path):
    claims = [claim("c1", "SYNTHETIC students were killed at the campus", "2024-07-16")]
    arts = [art("u1", "campus", "SYNTHETIC students were killed at the campus " * 5, "2024-07-16")]
    m.match(tmp_path, write(tmp_path, claims, arts), log=lambda s: None)
    ws = load_workbook(tmp_path / "matches" / "candidates.xlsx")["Candidates"]
    head = [c.value for c in ws[1]]
    assert "verdict" not in head and "stance" in head and ws.max_row == 2


def test_the_same_story_under_two_addresses_is_listed_once(tmp_path):
    claims = [claim("c1", "students were killed at the campus", "2024-07-16")]
    text = "SYNTHETIC students were killed at the campus " * 5
    arts = [
        art("u1", "Campus deaths | Paper", text, "2024-07-16"),
        art("u1%20", "Campus deaths", text, "2024-07-16"),
        art("u2", "Another story on the campus", text, "2024-07-16"),
    ]
    res = m.match(tmp_path, write(tmp_path, claims, arts), log=lambda s: None)
    assert len(res) == 2 and {r["url"] for r in res} >= {"u2"}


def test_candidates_are_balanced_by_language_and_outlet():
    arts = [
        {"outlet": "big", "title": f"English story {n}", "text": "SYNTHETIC text"} for n in range(8)
    ]
    arts += [{"outlet": f"bn{n}", "title": f"বাংলা খবর {n}", "text": "কৃত্রিম লেখা"} for n in range(4)]
    ranked = [(i, 1.0 - i / 100) for i in range(len(arts))]
    chosen = [i for i, _ in m.pick(arts, ranked, 6)]
    assert sum(1 for i in chosen if arts[i]["outlet"] == "big") == 3
    assert sum(1 for i in chosen if m.is_bangla(arts[i])) == 3
    assert chosen == sorted(chosen)  # still in rank order
