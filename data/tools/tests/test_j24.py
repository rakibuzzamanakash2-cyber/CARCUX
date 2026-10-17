"""CARCUX-J24 indexer. Every post below is SYNTHETIC: invented for the tests, with
example.org-style links where no real outlet is needed. No real site is contacted."""

import json
from pathlib import Path
from urllib.parse import parse_qs, urlsplit

import pytest
from openpyxl import load_workbook

from carcux_data.j24.__main__ import main
from carcux_data.j24.extract import classify, domain_of, load_outlets, verdict_words
from carcux_data.j24.fetch import Fetcher, Response
from carcux_data.j24.index import run
from carcux_data.j24.recurrence import summarise
from carcux_data.j24.sheet import check_sheet, write_sheet
from carcux_data.j24.sites import by_key

SALT = "s" * 40
OUTLETS = load_outlets()

# SYNTHETIC posts, shaped like WordPress REST API output.
POSTS = [
    {
        "id": 101, "date_gmt": "2024-07-20T10:00:00", "link": "https://rumorscanner.com/fc/101",
        "title": {"rendered": "SYNTHETIC: কোটা আন্দোলনে নিহতের সংখ্যা নিয়ে দাবি"},
        "categories": [7],
        "content": {"rendered": (
            "<p>Synthetic test text about a protest claim.</p>"
            '<a href="https://www.prothomalo.com/a/1">Prothom Alo</a>'
            '<a href="https://www.thedailystar.net/b/2">Daily Star</a>'
            '<a href="https://www.bbc.com/bengali/c/3">BBC</a>'
            '<a href="https://www.facebook.com/somepage/posts/999">post</a>'
            '<a href="https://perma.cc/ABCD-1234">archive</a>'
            '<a href="https://rumorscanner.com/fc/90">earlier check</a>'
            '<a href="https://unknown-portal.example/news/4">portal</a>'
            "<p>সিদ্ধান্ত: দাবিটি মিথ্যা।</p>"
        )},
    },
    {
        "id": 102, "date_gmt": "2024-07-25T08:00:00", "link": "https://rumorscanner.com/fc/102",
        "title": {"rendered": "SYNTHETIC: cricket score claim"},
        "categories": [8],
        "content": {"rendered": '<p>Unrelated sports item.</p><a href="https://www.prothomalo.com/x">x</a>'},
    },
]  # fmt: skip
CATEGORIES = [{"id": 7, "name": "গুজব"}, {"id": 8, "name": "খেলা"}]


class FakeWeb:
    """Serves the synthetic posts as a WordPress site would, and counts requests."""

    def __init__(self, robots="User-agent: *\nAllow: /\n", wordpress=True):
        self.robots = robots
        self.wordpress = wordpress
        self.calls: list[str] = []

    def __call__(self, url, headers):
        self.calls.append(url)
        assert "CARCUX-J24" in headers["User-Agent"]
        parts = urlsplit(url)
        if parts.path.endswith("robots.txt"):
            return Response(200, {}, self.robots.encode())
        if not self.wordpress:
            return Response(404, {}, b"<html>not here</html>")
        q = parse_qs(parts.query)
        if parts.path.endswith("/categories"):
            return Response(200, {}, json.dumps(CATEGORIES).encode())
        page = int(q["page"][0])
        if page > 1:
            return Response(400, {}, b'{"code":"rest_post_invalid_page_number"}')
        return Response(200, {"x-wp-totalpages": "1"}, json.dumps(POSTS).encode())


def fetcher(tmp_path, web):
    return Fetcher(tmp_path / "cache", opener=web, sleep=lambda s: None, log=lambda m: None)


def test_links_are_sorted_by_kind():
    assert domain_of("https://m.facebook.com/x") == "facebook.com"
    assert classify("https://www.prothomalo.com/a", OUTLETS).outlet.name == "Prothom Alo"
    assert classify("https://www.bbc.com/bengali/x", OUTLETS).outlet.country == "GB"
    assert classify("https://fb.watch/abc", OUTLETS).platform == "facebook"
    assert classify("https://archive.ph/xyz", OUTLETS).kind == "archive"
    assert classify("https://dghs.gov.bd/notice", OUTLETS).kind == "public_record"
    assert classify("https://factcheckbangla.afp.com/x", OUTLETS).kind == "fact_checker"
    # An archived copy of a news page counts as that outlet.
    wayback = "https://web.archive.org/web/20240720/https://www.thedailystar.net/news/1"
    assert classify(wayback, OUTLETS).outlet.name == "The Daily Star"
    assert classify("mailto:a@b.c", OUTLETS) is None


@pytest.mark.parametrize(
    ("title", "text", "expected"),
    [
        ("দাবিটি মিথ্যা", "", "false"),
        ("", "… সিদ্ধান্ত: ভিডিওটি এডিটেড", "altered"),
        ("Partly false: claim about X", "", "partly_false"),
        ("Misleading video", "", "misleading"),
        ("A claim about rain", "nothing conclusive here", None),
    ],
)
def test_verdict_words_are_a_hint(title, text, expected):
    assert verdict_words(title, text, [])[0] == expected


def test_index_writes_safe_public_records_and_restricted_links(tmp_path):
    web = FakeWeb()
    result = run(tmp_path, [by_key("rumorscanner-bn")], SALT, fetcher=fetcher(tmp_path, web),
                 log=lambda m: None)  # fmt: skip
    assert result.posts == 2
    rows = [json.loads(x) for x in (tmp_path / "factchecks.jsonl").read_text("utf-8").splitlines()]
    first = next(r for r in rows if r["factcheck_id"] == "rumorscanner-bn:101")
    assert first["verdict_hint"] == "false" and first["scope_hint"] is True
    assert first["categories"] == ["গুজব"]
    cited = {c["domain"]: c for c in first["cited"]}
    assert {"prothomalo.com", "thedailystar.net", "bbc.com", "perma.cc",
            "unknown-portal.example"} == set(cited)  # fmt: skip
    assert cited["unknown-portal.example"]["kind"] == "other"
    # The fact-checker's own links are not sources; social links are not in public files.
    public_text = (tmp_path / "factchecks.jsonl").read_text("utf-8")
    assert "facebook.com" not in public_text and "rumorscanner.com/fc/90" not in public_text
    assert first["social_links"][0]["platform"] == "facebook"
    assert len(first["social_links"][0]["url_ref"]) == 64
    restricted = (tmp_path / "restricted" / "social_links.jsonl").read_text("utf-8")
    assert "facebook.com/somepage/posts/999" in restricted
    second = next(r for r in rows if r["factcheck_id"] == "rumorscanner-bn:102")
    assert second["scope_hint"] is False
    # No article body is kept.
    assert "Synthetic test text" not in public_text


def test_second_run_uses_the_cache(tmp_path):
    web = FakeWeb()
    run(tmp_path, [by_key("rumorscanner-bn")], SALT, fetcher=fetcher(tmp_path, web),
        log=lambda m: None)  # fmt: skip
    before = len(web.calls)
    run(tmp_path, [by_key("rumorscanner-bn")], SALT, fetcher=fetcher(tmp_path, web),
        log=lambda m: None)  # fmt: skip
    assert len(web.calls) == before  # everything came from the cache


def test_robots_txt_is_respected(tmp_path):
    web = FakeWeb(robots="User-agent: *\nDisallow: /wp-json/\n")
    result = run(tmp_path, [by_key("rumorscanner-bn")], SALT, fetcher=fetcher(tmp_path, web),
                 log=lambda m: None)  # fmt: skip
    assert result.posts == 0 and "robots.txt" in result.sites[0].problem
    assert not any("/posts" in c for c in web.calls)


def test_a_site_that_is_not_wordpress_is_reported(tmp_path):
    web = FakeWeb(wordpress=False)
    result = run(tmp_path, [by_key("factwatch")], SALT, fetcher=fetcher(tmp_path, web),
                 log=lambda m: None)  # fmt: skip
    assert result.posts == 0 and "did not answer as WordPress" in result.sites[0].problem


def test_requests_are_spaced_out(tmp_path):
    pauses = []
    f = Fetcher(tmp_path / "c", opener=FakeWeb(), sleep=pauses.append, delay=2.5)
    run(tmp_path, [by_key("rumorscanner-bn")], SALT, fetcher=f, log=lambda m: None)
    assert pauses and set(pauses) == {2.5}


def _records(tmp_path):
    run(tmp_path, [by_key("rumorscanner-bn")], SALT, fetcher=fetcher(tmp_path, FakeWeb()),
        log=lambda m: None)  # fmt: skip
    path = tmp_path / "factchecks.jsonl"
    return [json.loads(x) for x in path.read_text("utf-8").splitlines()]


def test_recurrence_counts_outlets_per_fact_check(tmp_path):
    s = summarise(_records(tmp_path))
    assert s["factchecks"] == 1 and s["factchecks_all"] == 2
    assert {t["domain"] for t in s["top_sources"]} == {"prothomalo.com", "thedailystar.net",
                                                       "bbc.com"}  # fmt: skip
    assert s["unregistered_domains"] == [("unknown-portal.example", 1)]
    assert s["recurring_sources"] == 0


def test_claim_sheet_round_trip(tmp_path):
    path = tmp_path / "claim_sheet.xlsx"
    assert write_sheet(_records(tmp_path), path) == 2
    wb = load_workbook(path)
    assert wb.sheetnames == ["How to fill", "Claims", "Lists"]
    ws = wb["Claims"]
    header = [c.value for c in ws[1]]
    assert ws.cell(row=2, column=1).value == "rumorscanner-bn:101"  # likely in scope first
    # The person's columns start empty.
    for name in ("claim", "verdict", "checked_by"):
        assert ws.cell(row=2, column=header.index(name) + 1).value is None
    assert check_sheet(path) == []  # nothing marked in scope yet
    col = {n: header.index(n) + 1 for n in header}
    ws.cell(row=2, column=col["in_scope"], value="yes")
    ws.cell(row=2, column=col["verdict"], value="probably")
    wb.save(path)
    problems = "\n".join(check_sheet(path))
    assert "claim is empty" in problems and "verdict 'probably'" in problems
    for name, value in [("claim", "Deaths at a place on a day (synthetic)"),
                        ("claim_type", "count"), ("event_type", "armed_violence"),
                        ("event_date", "2024-07-19"), ("verdict", "false"),
                        ("checked_by", "annotator-x")]:  # fmt: skip
        ws.cell(row=2, column=col[name], value=value)
    wb.save(path)
    assert check_sheet(path) == []


def test_cli_needs_a_salt(tmp_path, monkeypatch, capsys):
    monkeypatch.delenv("CARCUX_J24_SALT", raising=False)
    assert main(["index", "--out", str(tmp_path)]) == 2
    assert "CARCUX_J24_SALT" in capsys.readouterr().err


def test_cli_will_not_overwrite_a_sheet(tmp_path):
    _records(tmp_path)
    assert main(["claim-sheet", "--out", str(tmp_path)]) == 0
    assert main(["claim-sheet", "--out", str(tmp_path)]) == 2
    assert Path(tmp_path / "claim_sheet.xlsx").exists()
