"""Claim-sheet batches for a helper. Every fact-check below is SYNTHETIC."""

from openpyxl import load_workbook

from carcux_data.j24.__main__ import main
from carcux_data.j24.batches import HINT, merge_sheets, read_rows, split_sheet, suggest_same
from carcux_data.j24.sheet import check_sheet, code, write_sheet


def rec(i, title, day, likely=True):
    return {"factcheck_id": f"fc:{i}", "fact_checker": "synthetic", "url": f"https://fc.example/{i}",
            "published_at": f"2024-07-{day:02d}T10:00:00", "title": title,
            "verdict_hint": "false", "scope_hint": likely}  # fmt: skip


RECORDS = [
    rec(1, "SYNTHETIC: twenty killed at market in District X video claim", 20),
    rec(2, "SYNTHETIC: video claim twenty killed at market in District X false", 21),
    rec(3, "SYNTHETIC: curfew extended in Town Y rumour", 22),
    rec(4, "SYNTHETIC: cricket player quote", 10, likely=False),
] + [rec(10 + k, f"SYNTHETIC: unrelated topic number {k} alpha{k} beta{k}", 1 + k % 28)
     for k in range(30)]  # fmt: skip


def sheet(tmp_path):
    path = tmp_path / "claim_sheet.xlsx"
    write_sheet(RECORDS, path)
    return path


def test_code_reads_bilingual_choices():
    assert code("false — মিথ্যা") == "false" and code(" yes ") == "yes"


def test_similar_titles_a_day_apart_are_suggested_as_the_same_claim():
    hints = suggest_same(RECORDS)
    assert hints == {"fc:2": "fc:1"}


def test_split_keeps_suggested_pairs_together_and_carries_work_over(tmp_path):
    path = sheet(tmp_path)
    wb = load_workbook(path)
    ws = wb["Claims"]
    header = [c.value for c in ws[1]]
    for row in ws.iter_rows(min_row=2):
        if row[0].value == "fc:3":
            row[header.index("claim")].value = "Curfew extended in Town Y (synthetic)"
    wb.save(path)

    written = split_sheet(path, tmp_path / "b", size=2, checked_by="helper-1")
    rows = [read_rows(p) for p, _, _ in written]
    together = next(b for b in rows if any(r["factcheck_id"] == "fc:1" for r in b))
    assert {r["factcheck_id"] for r in together} == {"fc:1", "fc:2"}
    assert next(r for r in together if r["factcheck_id"] == "fc:2")[HINT] == "fc:1"
    flat = {r["factcheck_id"]: r for b in rows for r in b}
    assert len(flat) == len(RECORDS)
    assert flat["fc:3"]["claim"] == "Curfew extended in Town Y (synthetic)"
    assert flat["fc:4"]["checked_by"] == "helper-1"

    wb = load_workbook(written[0][0])
    assert wb.sheetnames == ["Start here", "Claims", "Lists"]
    assert str(wb["Start here"]["B1"].value).startswith("=SUMPRODUCT")
    assert wb["Claims"]["D2"].hyperlink is not None
    assert wb["Lists"]["D2"].value.startswith("false — ")


def test_bilingual_choices_merge_back_to_codes_and_check(tmp_path):
    path = sheet(tmp_path)
    written = split_sheet(path, tmp_path / "b", size=40)
    wb = load_workbook(written[0][0])
    ws = wb["Claims"]
    h = [c.value for c in ws[1]]
    row = next(r for r in ws.iter_rows(min_row=2) if r[0].value == "fc:1")
    for name, value in [("in_scope", "yes — হ্যাঁ"), ("claim", "Twenty killed (synthetic)"),
                        ("claim_type", "count — সংখ্যা"), ("event_type", "armed_violence — x"),
                        ("event_date", "2024-07-20"), ("verdict", "false — মিথ্যা"),
                        ("checked_by", "helper-1")]:  # fmt: skip
        row[h.index(name)].value = value
    wb.save(written[0][0])
    assert check_sheet(written[0][0]) == []

    out = tmp_path / "merged.xlsx"
    assert merge_sheets([p for p, _, _ in written], out) == len(RECORDS)
    merged = {r["factcheck_id"]: r for r in read_rows(out)}
    assert merged["fc:1"]["verdict"] == "false" and merged["fc:1"]["in_scope"] == "yes"
    assert check_sheet(out) == []


def test_cli_split_refuses_to_overwrite_batches_and_merge_reports(tmp_path, capsys):
    sheet(tmp_path)
    assert main(["split-sheet", "--out", str(tmp_path), "--size", "10"]) == 0
    assert "with a maybe-same hint" in capsys.readouterr().out
    assert main(["split-sheet", "--out", str(tmp_path)]) == 2
    assert main(["merge-sheets", "--out", str(tmp_path)]) == 0
    assert "complete" in capsys.readouterr().out
    assert main(["merge-sheets", "--out", str(tmp_path)]) == 2  # will not overwrite
