"""Split the claim sheet into small, friendlier batches for a helper, and merge them back.

A batch is the claim sheet's "Claims" table for about 25 fact-checks, with:

- a "Start here" page in Bangla and English, a progress counter and three worked examples
  (SYNTHETIC: made-up places and numbers);
- dropdowns that show each choice's meaning in Bangla ("false — মিথ্যা"); only the English
  code before " — " counts, so checking and merging read the same values as before;
- clickable links, wrapped titles, and rows that turn green when complete;
- a `maybe_same_as` hint (grey, filled by the tool): another row whose title shares most of
  its words and was published within a few days. It is only a hint; the person decides and
  copies it into `same_claim_as` if it really is the same claim. Rows that may be the same
  claim are kept in the same batch.

Nothing in the person's columns is filled by the tool, except `checked_by` when the person
running the split gives a pseudonym for the helper. Work already in the source sheet is
carried over, so splitting a half-filled sheet loses nothing.
"""

from __future__ import annotations

import re
from collections import Counter
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from carcux_data.j24.sheet import AUTO, CHOICES, PERSON, code

HINT = "maybe_same_as"
BATCH_COLUMNS = AUTO + [HINT] + PERSON

BANGLA = {
    "yes": "হ্যাঁ, আন্দোলন/পরবর্তী ঘটনা",
    "no": "না, অন্য বিষয়",
    "occurrence": "ঘটনা ঘটেছে কি না",
    "count": "সংখ্যা (নিহত, আহত, গ্রেপ্তার)",
    "attribution": "কে কী বলেছে",
    "media_authenticity": "ছবি/ভিডিও আসল কি না",
    "official_action": "সরকারি পদক্ষেপ (কারফিউ, আদেশ)",
    "protest": "বিক্ষোভ, মিছিল",
    "armed_violence": "গুলি, সশস্ত্র সংঘর্ষ",
    "mob_violence": "গণপিটুনি, মব হামলা",
    "arson": "অগ্নিসংযোগ",
    "attack_on_place_of_worship": "উপাসনালয়ে হামলা",
    "communications_shutdown": "ইন্টারনেট/যোগাযোগ বন্ধ",
    "curfew": "কারফিউ",
    "mass_arrest": "গণগ্রেপ্তার",
    "publication": "কোনো খবর/বই/পোস্ট প্রকাশ",
    "official_statement": "সরকারি বা দাপ্তরিক বক্তব্য",
    "other": "অন্যান্য",
    "false": "মিথ্যা",
    "misleading": "বিভ্রান্তিকর",
    "altered": "বিকৃত / এডিট করা",
    "partly_false": "আংশিক মিথ্যা",
    "partly_true": "আংশিক সত্য",
    "true": "সত্য",
    "unproven": "প্রমাণিত নয়",
    "B": "প্রাথমিক প্রমাণ (রেকর্ড, মূল ভিডিও, মূল সংবাদমাধ্যম)",
    "C": "শুধু অন্য প্রতিবেদন",
}
_ALL_CODES = {v for values in CHOICES.values() for v in values}
assert _ALL_CODES <= set(BANGLA), sorted(_ALL_CODES - set(BANGLA))


def shown(value: str) -> str:
    return f"{value} — {BANGLA[value]}"


START = [
    ("এটা কী? / What is this?",
     "প্রতিটি সারি একটি ফ্যাক্ট-চেক। লিংকে ক্লিক করে পড়ুন, তারপর সাদা ঘরগুলো পূরণ করুন। "
     "ধূসর ঘর টুলটি পূরণ করেছে, বদলাবেন না।\n"
     "Each row is one fact-check. Click the link, read it, fill the white cells. Grey cells "
     "were filled by the tool: leave them."),
    ("৩টি ধাপ / 3 steps",
     "১. in_scope: আন্দোলন বা ১৫ আগস্ট পর্যন্ত পরবর্তী ঘটনা নিয়ে? না হলে 'no' দিয়ে পরের সারিতে যান "
     "(১০ সেকেন্ড)।\n"
     "২. maybe_same_as-এ কিছু লেখা থাকলে দুটো মিলিয়ে দেখুন। একই দাবি হলে সেটা same_claim_as-এ কপি "
     "করে পরের সারিতে যান।\n"
     "৩. বাকিগুলোতে: claim (এক বাক্যে, নিরপেক্ষভাবে), তারপর ড্রপডাউনগুলো, place, event_date, verdict।\n"
     "1. in_scope: about the uprising or its aftermath (to 15 August)? If not, choose 'no' "
     "and move on. 2. If maybe_same_as shows an id and it is the same claim, copy it into "
     "same_claim_as and move on. 3. Otherwise write the claim and pick from the dropdowns."),
    ("claim লেখার নিয়ম / Writing the claim",
     "কী, কোথায়, কখন — এক বাক্যে, যা দাবি করা হয়েছিল সেটাই (রায় নয়)। কোনো সাধারণ মানুষের নাম নয়: "
     "'একজন শিক্ষার্থী', 'একজন ব্যক্তি' লিখুন। সংগঠন বা প্রতিষ্ঠানের নাম লেখা যাবে।\n"
     "What, where, when, in one sentence: the claim as it was made, not the verdict. No names of "
     "private people. Bangla or English, whichever is quicker."),
    ("verdict",
     "ফ্যাক্ট-চেকার যে রায় দিয়েছে, সেটাই বেছে নিন — আপনার নিজের মত নয়। রায় না থাকলে খালি রাখুন। "
     "verdict_hint-এ টুলটি লেখার শব্দ দেখিয়েছে, তবু লেখাটি পড়ে মিলিয়ে নিন।\n"
     "The fact-checker's verdict, never your own view. verdict_hint shows words found in the "
     "article; confirm them by reading."),
    ("evidence_tier",
     "B: ফ্যাক্ট-চেকটি প্রাথমিক প্রমাণে দাঁড়িয়ে (হাসপাতাল বা সরকারি রেকর্ড, মূল ছবি/ভিডিও, যাকে "
     "উদ্ধৃত করা হয়েছে সেই সংবাদমাধ্যম)। C: শুধু অন্য প্রতিবেদনের ওপর। নিশ্চিত না হলে notes-এ লিখুন।"),
    ("checked_by", "আপনার ছদ্মনাম (pseudonym), আসল নাম নয়। / Your pseudonym, never your name."),
    ("গোপনীয়তা / Privacy",
     "কোনো পোস্টের লেখা, মানুষের নাম, অ্যাকাউন্টের নাম বা ফোন নম্বর শিটে কপি করবেন না।\n"
     "Do not copy post text, names, handles or phone numbers into the sheet."),
    ("বিরতি / Breaks",
     "কিছু লেখায় মৃত্যু ও সহিংসতার বর্ণনা আছে। ভিডিও দেখার দরকার নেই। যেকোনো সারি বাদ দিতে পারেন "
     "(notes-এ 'skipped' লিখুন), যখন খুশি থামতে পারেন। একটি ফাইল ≈ ৪০–৬০ মিনিট।\n"
     "Some articles describe deaths and violence. No need to watch videos. Skip any row "
     "(write 'skipped' in notes) and stop whenever you like. One file ≈ 40–60 minutes."),
]  # fmt: skip

# SYNTHETIC worked examples: "District X", "Town Y" and the numbers are made up.
EXAMPLES = [
    ("A fact-check says a viral post claimed 20 people were killed at a market in District X "
     "on 20 July; hospital records showed 3. Verdict: false.",
     ["yes", "", "Twenty people were killed at a market in District X on 20 July 2024", "count",
      "armed_violence", "District X", "2024-07-20", "false", "B"]),
    ("A fact-check says a video shared as 'a police station burning in Town Y on 5 August' is "
     "from 2019. Verdict: misleading.",
     ["yes", "", "A video shows a police station burning in Town Y on 5 August 2024",
      "media_authenticity", "arson", "Town Y", "2024-08-05", "misleading", "B"]),
    ("A fact-check about a cricket player's quote, published 10 July.",
     ["no", "", "", "", "", "", "", "", ""]),
]  # fmt: skip


def _tokens(title: str) -> set[str]:
    return {t for t in re.split(r"[\s\"'‘’“”:;,.!?।()\[\]{}|/—–-]+", title.lower()) if len(t) > 1}


def _day(text: str | None) -> date | None:
    try:
        return date.fromisoformat(str(text)[:10]) if text else None
    except ValueError:
        return None


def suggest_same(rows: list[dict], *, overlap: float = 0.5, days: int = 3) -> dict[str, str]:
    """{factcheck_id: earlier factcheck_id} for rows that may check the same claim."""
    toks = {r["factcheck_id"]: _tokens(str(r.get("title") or "")) for r in rows}
    # Words in many titles ("fact", "check", "viral", "ভাইরাল") say nothing about the claim.
    df = Counter(t for ts in toks.values() for t in ts)
    common = {t for t, n in df.items() if n > max(3, len(rows) // 10)}
    out: dict[str, str] = {}
    ordered = sorted(rows, key=lambda r: (str(r.get("published_at") or ""), r["factcheck_id"]))
    for i, r in enumerate(ordered):
        a, da = toks[r["factcheck_id"]] - common, _day(r.get("published_at"))
        if len(a) < 3 or da is None:
            continue
        for e in ordered[:i]:
            b, db = toks[e["factcheck_id"]] - common, _day(e.get("published_at"))
            if db is None or abs((da - db).days) > days or len(b) < 3:
                continue
            if len(a & b) / len(a | b) >= overlap:
                out[r["factcheck_id"]] = out.get(e["factcheck_id"], e["factcheck_id"])
                break
    return out


def read_rows(path: Path) -> list[dict]:
    """Every row of a claim sheet or batch, as {column: value}, English codes only."""
    ws = load_workbook(path, read_only=True)["Claims"]
    it = ws.iter_rows(values_only=True)
    header = [str(h) if h is not None else "" for h in next(it)]
    rows = []
    for values in it:
        r = {h: v for h, v in zip(header, values, strict=False) if h}
        if not r.get("factcheck_id"):
            continue
        for name in CHOICES:
            if r.get(name) is not None:
                r[name] = code(r[name])
        rows.append(r)
    return rows


def _write(rows: list[dict], path: Path, *, friendly: bool) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Claims"
    columns = BATCH_COLUMNS
    ws.append(columns)
    letter = {n: ws.cell(row=1, column=i).column_letter for i, n in enumerate(columns, 1)}
    grey, head = PatternFill("solid", fgColor="E7E6E6"), PatternFill("solid", fgColor="1F5F3A")
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), head
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for r in rows:
        values = []
        for n in columns:
            v = r.get(n)
            if v is None:
                v = ""
            if friendly and n in CHOICES and v in BANGLA:
                v = shown(v)
            values.append(v)
        ws.append(values)
    last = max(2, len(rows) + 1)
    for row in ws.iter_rows(min_row=2, max_row=last):
        for c in row:
            name = columns[c.column - 1]
            c.alignment = Alignment(wrap_text=name in ("title", "claim", "notes"), vertical="top")
            if name in AUTO or name == HINT:
                c.fill = grey
            if name == "url" and c.value:
                c.hyperlink = str(c.value)
                c.style = "Hyperlink"
                c.fill = grey
    widths = {"factcheck_id": 20, "url": 28, "title": 55, "claim": 55, "notes": 28,
              HINT: 20, "same_claim_as": 20, "event_type": 24, "claim_type": 22,
              "verdict": 18}  # fmt: skip
    for n in columns:
        ws.column_dimensions[letter[n]].width = widths.get(n, 14)
    ws.freeze_panes = "B2"

    lists = wb.create_sheet("Lists")
    for j, (name, values) in enumerate(CHOICES.items(), start=1):
        col = lists.cell(row=1, column=j).column_letter
        lists.cell(row=1, column=j, value=name).font = Font(bold=True)
        for k, v in enumerate(values, start=2):
            lists.cell(row=k, column=j, value=shown(v) if friendly else v)
        lists.column_dimensions[col].width = 36
        dv = DataValidation(type="list", formula1=f"=Lists!${col}$2:${col}${len(values) + 1}",
                            allow_blank=True)  # fmt: skip
        dv.add(f"{letter[name]}2:{letter[name]}{last}")
        ws.add_data_validation(dv)

    if friendly:
        _friendly(wb, ws, letter, last, len(rows))
        wb.active = 1
    wb.save(path)


def _done_test(letter: dict[str, str], r: str) -> str:
    """An Excel test that row r is complete: out of scope, a duplicate, or filled."""
    L = {n: f"{letter[n]}{r}" for n in letter}  # noqa: N806
    filled = "*".join(f'({L[n]}<>"")' for n in ("claim", "claim_type", "event_type",
                                                  "event_date", "checked_by"))  # fmt: skip
    return (f'((LEFT({L["in_scope"]},2)="no")+(LEFT({L["in_scope"]},3)="yes")'
            f'*((({L["same_claim_as"]}<>"")+{filled})>0))>0')  # fmt: skip


def _friendly(wb: Workbook, ws, letter: dict[str, str], last: int, n: int) -> None:
    green = PatternFill("solid", fgColor="D8EFD9")
    one = _done_test({k: f"${v}" for k, v in letter.items()}, "2")
    ws.conditional_formatting.add(
        f"{letter['in_scope']}2:{letter['notes']}{last}",
        FormulaRule(formula=[one], fill=green, stopIfTrue=False),
    )
    ws.conditional_formatting.add(
        f"{letter[HINT]}2:{letter[HINT]}{last}",
        FormulaRule(formula=[f'${letter[HINT]}2<>""'], fill=PatternFill("solid", fgColor="FFE8A3")),
    )

    start = wb.create_sheet("Start here", 0)
    start.column_dimensions["A"].width = 34
    start.column_dimensions["B"].width = 100
    start["A1"] = "অগ্রগতি / Progress"
    start["A1"].font = Font(bold=True, size=14)
    rng = {k: f"Claims!{v}2:{v}{last}" for k, v in letter.items()}
    filled = "*".join(f'({rng[k]}<>"")' for k in ("claim", "claim_type", "event_type",
                                                    "event_date", "checked_by"))  # fmt: skip
    count = (f'SUMPRODUCT(--(((LEFT({rng["in_scope"]},2)="no")+(LEFT({rng["in_scope"]},3)="yes")'
             f'*((({rng["same_claim_as"]}<>"")+{filled})>0))>0))')  # fmt: skip
    start["B1"] = f'={count}&" of {n} done  /  {n} টির মধ্যে "&{count}&" টি শেষ"'
    start["B1"].font = Font(bold=True, size=14, color="1F5F3A")
    start.append([])
    for a, b in START:
        start.append([a, b])
    start.append([])
    start.append(["উদাহরণ / Examples", "Made-up examples (SYNTHETIC) — copy the pattern."])
    head = ["in_scope", "same_claim_as", "claim", "claim_type", "event_type", "place",
            "event_date", "verdict", "evidence_tier"]  # fmt: skip
    for text, values in EXAMPLES:
        start.append([])
        start.append(["The fact-check", text])
        for h, v in zip(head, values, strict=True):
            if v:
                start.append([f"   {h}", shown(v) if v in BANGLA else v])
    for row in start.iter_rows(min_row=2):
        row[0].font = Font(bold=True)
        row[0].alignment = Alignment(vertical="top")
        if len(row) > 1:
            row[1].alignment = Alignment(wrap_text=True, vertical="top")


def split_sheet(source: Path, out_dir: Path, *, size: int = 25, only_likely: bool = False,
                checked_by: str = "") -> list[tuple[Path, int, int]]:  # fmt: skip
    """Write batch_01.xlsx, batch_02.xlsx, … to out_dir; [(path, rows, rows_with_hint)]."""
    rows = read_rows(source)
    if only_likely:
        rows = [r for r in rows if r.get("scope_hint") or r.get("in_scope")]
    hints = suggest_same(rows)
    for r in rows:
        r[HINT] = hints.get(r["factcheck_id"], "")
        if checked_by and not r.get("checked_by"):
            r["checked_by"] = checked_by
    # Keep the source order (likely first, then by date), but pull each row that may repeat
    # another next to it, so both land in the same batch.
    by_id = {r["factcheck_id"]: r for r in rows}
    followers: dict[str, list[dict]] = {}
    for r in rows:
        if r[HINT] in by_id:
            followers.setdefault(r[HINT], []).append(r)
    groups = [[r, *followers.get(r["factcheck_id"], [])] for r in rows
              if r[HINT] not in by_id]  # fmt: skip
    batches: list[list[dict]] = [[]]
    for g in groups:
        if batches[-1] and len(batches[-1]) + len(g) > size:
            batches.append([])
        batches[-1].extend(g)
    out_dir.mkdir(parents=True, exist_ok=True)
    written = []
    for i, b in enumerate(b for b in batches if b):
        path = out_dir / f"batch_{i + 1:02d}.xlsx"
        _write(b, path, friendly=True)
        written.append((path, len(b), sum(1 for r in b if r[HINT])))
    return written


def merge_sheets(files: list[Path], out: Path) -> int:
    """Join batches (in any order) into one claim sheet with English codes."""
    seen: dict[str, Path] = {}
    rows = []
    for f in sorted(files):
        for r in read_rows(f):
            fid = r["factcheck_id"]
            if fid in seen:
                raise ValueError(f"{fid} is in both {seen[fid].name} and {f.name}")
            seen[fid] = f
            rows.append(r)
    _write(rows, out, friendly=False)
    return len(rows)
