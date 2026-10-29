"""The automatic stance baseline, and its evaluation against the human labels.

For each candidate pair (claim, article) from the matcher, `score` finds the article
sentences closest to the claim and asks a multilingual natural-language-inference model
whether each sentence entails, contradicts or is neutral to the claim. It also counts cue
words that signal a correction ("ভুয়া", "fake", "fact-check") or doubt ("অস্বীকার",
"denied", "unverified"). The claim is put to the model in the article's language (the
Bangla translation for Bangla articles).

The result is a table of features per pair. `predict` turns features into one of the five
labels with a few thresholds. The thresholds are chosen only on the development split of
the human labels (a fixed 20% of claims) and reported on the remaining 80%: the protocol
does not allow tuning on the test set.

`evaluate` reads the returned labelling sheets, measures agreement between the labellers
on the shared pairs (Krippendorff's alpha), builds the gold labels, and scores the model
and two simple baselines on the test split.
"""

from __future__ import annotations

import itertools
import json
import random
import re
from collections import Counter, defaultdict
from collections.abc import Callable
from pathlib import Path

from carcux_data.j24.match import is_bangla, load_texts, tokens

LABELS = ("repeats", "questions", "corrects", "mentions", "unrelated")
NLI_MODEL = "MoritzLaurer/mDeBERTa-v3-base-xnli-multilingual-nli-2mil7"
EMBED_MODEL = "intfloat/multilingual-e5-small"
SENTENCES = 8  # sentences per article kept after a word-overlap filter
TOP = 3  # of those, the most similar to the claim, put to the NLI model
DEV_SHARE, SEED = 0.2, 2024

CORRECT_CUES = ("ভুয়া", "ভূয়া", "গুজব", "মিথ্যা", "ভিত্তিহীন", "বানোয়াট", "সত্য নয়", "ফ্যাক্টচেক",
                "fake", "false", "rumour", "rumor", "fact-check", "fact check", "debunk",
                "baseless", "fabricated")  # fmt: skip
DOUBT_CUES = ("অস্বীকার", "প্রতিবাদ", "নিশ্চিত হওয়া যায়নি", "নিশ্চিত নয়", "অভিযোগ", "দাবি করেছ",
              "denied", "deny", "denies", "unverified", "could not be verified", "alleged",
              "claimed", "unconfirmed", "disputed", "rejected")  # fmt: skip


def sentences(text: str) -> list[str]:
    parts = re.split(r"(?<=[।.!?])\s+", " ".join(text.split()))
    return [p for p in parts if 25 <= len(p) <= 600]


def _cues(text: str, cues: tuple[str, ...]) -> int:
    low = text.lower()
    return sum(low.count(c.lower()) for c in cues)


def candidate_sentences(claim: str, article: dict, k: int = SENTENCES) -> list[str]:
    """The title and the k sentences sharing most words with the claim."""
    want = set(tokens(claim))
    sents = sentences(article.get("text") or "")
    ranked = sorted(sents, key=lambda s: -len(want & set(tokens(s))))[:k]
    title = (article.get("title") or "").strip()
    return ([title] if title else []) + ranked


# ------------------------------------------------------------------ scoring with models


def _models(log):  # pragma: no cover - needs the models, run on the user's machine
    import numpy as np
    from sentence_transformers import SentenceTransformer
    from transformers import AutoModelForSequenceClassification, AutoTokenizer, pipeline

    log(f"  loading {EMBED_MODEL} and {NLI_MODEL}")
    emb = SentenceTransformer(EMBED_MODEL, device="cpu")
    tok = AutoTokenizer.from_pretrained(NLI_MODEL)
    mdl = AutoModelForSequenceClassification.from_pretrained(NLI_MODEL)
    nli = pipeline("text-classification", model=mdl, tokenizer=tok, device=-1, top_k=None,
                   truncation=True, max_length=256)  # fmt: skip

    def embed(claim: str, sents: list[str]) -> list[float]:
        q = emb.encode(["query: " + claim], normalize_embeddings=True)
        p = emb.encode(["passage: " + s for s in sents], normalize_embeddings=True)
        return list((np.asarray(p) @ np.asarray(q).T).ravel())

    def infer(pairs: list[tuple[str, str]]) -> list[dict[str, float]]:
        res = nli([{"text": p, "text_pair": h} for p, h in pairs], batch_size=8)
        return [{d["label"].lower(): d["score"] for d in r} for r in res]

    return embed, infer


def score(out: Path, claims_bn: Path | None = None, *, embed=None, infer=None,
          log: Callable[[str], None] = print) -> list[dict]:  # fmt: skip
    """Features for every candidate pair in matches/candidates.jsonl."""
    cands = [json.loads(x) for x in (out / "matches" / "candidates.jsonl")
             .read_text("utf-8").splitlines() if x.strip()]  # fmt: skip
    texts = {a["url"]: a for a in load_texts(out)}
    claims = {}
    path = out / "matches" / "candidates.xlsx"
    if path.exists():
        from openpyxl import load_workbook

        ws = load_workbook(path, read_only=True)["Candidates"]
        for row in ws.iter_rows(min_row=2, values_only=True):
            claims.setdefault(row[1], row[2])
    bn = {}
    if claims_bn and claims_bn.exists():
        for line in claims_bn.read_text("utf-8").splitlines():
            if line.strip():
                r = json.loads(line)
                bn[r["claim_group"]] = r["claim_bn"]
                if r.get("claim_en"):
                    claims[r["claim_group"]] = r["claim_en"]  # the current wording
    if embed is None or infer is None:
        embed, infer = _models(log)
    rows = []
    for n, c in enumerate(cands, 1):
        a = texts.get(c["url"], {"title": c.get("title"), "text": c.get("snippet")})
        claim_en = claims.get(c["claim_group"], "")
        bangla = is_bangla(a)
        claim = bn.get(c["claim_group"], claim_en) if bangla else claim_en
        sents = candidate_sentences(claim, a)
        sims = embed(claim, sents) if sents else []
        order = sorted(range(len(sents)), key=lambda i: -sims[i])[:TOP]
        probs = infer([(sents[i], claim) for i in order]) if order else []
        best_e = max((p.get("entailment", 0.0) for p in probs), default=0.0)
        best_c = max((p.get("contradiction", 0.0) for p in probs), default=0.0)
        whole = f"{a.get('title') or ''} {a.get('text') or ''}"
        rows.append({
            "pair_id": f"p{n:05d}", "claim_group": c["claim_group"], "url": c["url"],
            "bangla": bangla, "sim_max": round(float(max(sims, default=0.0)), 4),
            "entail": round(best_e, 4), "contra": round(best_c, 4),
            "correct_cues": _cues(whole, CORRECT_CUES), "doubt_cues": _cues(whole, DOUBT_CUES),
            "evidence": (sents[order[0]] if order else "")[:300],
        })  # fmt: skip
        if n % 200 == 0:
            log(f"  {n} of {len(cands)} pairs scored")
    target = out / "matches" / "stance_features.jsonl"
    target.write_text("".join(json.dumps(r, ensure_ascii=False) + "\n" for r in rows),
                      encoding="utf-8")  # fmt: skip
    return rows


# ------------------------------------------------------------------ from features to labels

DEFAULTS = {"sim": 0.80, "entail": 0.5, "contra": 0.5, "cues": 1, "strong": 10**6}


def predict(f: dict, p: dict = DEFAULTS) -> str:
    """One label from the features. The order follows the guide's decision steps."""
    if f["sim_max"] < p["sim"]:
        return "unrelated"
    if f["correct_cues"] >= p.get("strong", 10**6):
        return "corrects"  # many correction words: a fact-check or a denial story
    if f["contra"] >= p["contra"] and f["correct_cues"] >= p["cues"]:
        return "corrects"
    if f["doubt_cues"] >= p["cues"] and (
        f["contra"] >= p["contra"] / 2 or f["entail"] >= p["entail"]
    ):
        return "questions"
    if f["entail"] >= p["entail"]:
        return "repeats"
    if f["contra"] >= p["contra"]:
        return "corrects"
    return "mentions"


GRID = {"sim": [0.70, 0.78, 0.79, 0.80, 0.81, 0.82, 0.83, 0.84, 0.85, 0.86, 0.88],
        "entail": [0.3, 0.4, 0.5, 0.6, 0.7], "contra": [0.3, 0.4, 0.5, 0.6, 0.7, 0.8, 0.9],
        "cues": [1, 2, 3], "strong": [3, 5, 8, 10**6]}  # fmt: skip


def macro_f1(gold: list[str], pred: list[str]) -> float:
    scores = []
    for lab in LABELS:
        tp = sum(g == lab and p == lab for g, p in zip(gold, pred, strict=True))
        fp = sum(g != lab and p == lab for g, p in zip(gold, pred, strict=True))
        fn = sum(g == lab and p != lab for g, p in zip(gold, pred, strict=True))
        if tp + fp + fn:
            prec = tp / (tp + fp) if tp + fp else 0.0
            rec = tp / (tp + fn) if tp + fn else 0.0
            scores.append(2 * prec * rec / (prec + rec) if prec + rec else 0.0)
    return sum(scores) / len(scores) if scores else 0.0


def fit(features: list[dict], gold: list[str]) -> dict:
    """The thresholds with the best macro-F1 on the development pairs."""
    best, best_p = -1.0, DEFAULTS
    for values in itertools.product(*GRID.values()):
        p = dict(zip(GRID, values, strict=True))
        s = macro_f1(gold, [predict(f, p) for f in features])
        if s > best:
            best, best_p = s, p
    return best_p


# ------------------------------------------------------------------ the human labels

LABEL_OF = {lab: lab for lab in LABELS}


def _label(cell) -> str | None:
    if not cell:
        return None
    head = str(cell).split("—")[0].strip().lower()
    return LABEL_OF.get(head)


def read_sheets(paths: list[Path]) -> dict[str, list[tuple[str, str]]]:
    """pair_id -> [(annotator, label)], from the returned part sheets."""
    from openpyxl import load_workbook

    labels: dict[str, list[tuple[str, str]]] = defaultdict(list)
    for path in paths:
        ws = load_workbook(path, read_only=True).worksheets[0]
        head = None
        for row in ws.iter_rows(values_only=True):
            if head is None:
                head = list(row)
                continue
            r = dict(zip(head, row, strict=False))
            lab = _label(r.get("stance"))
            if r.get("pair_id") and lab:
                labels[r["pair_id"]].append((str(r.get("annotator") or path.stem), lab))
    return labels


def krippendorff_nominal(units: list[list[str]]) -> float | None:
    """Krippendorff's alpha for nominal labels; units with fewer than two labels are left out."""
    units = [u for u in units if len(u) >= 2]
    if not units:
        return None
    pairs = Counter()
    for u in units:
        m = len(u)
        for a, b in itertools.permutations(range(m), 2):
            pairs[(u[a], u[b])] += 1 / (m - 1)
    n = sum(pairs.values())
    totals = Counter()
    for (a, _), v in pairs.items():
        totals[a] += v
    d_o = sum(v for (a, b), v in pairs.items() if a != b) / n
    d_e = sum(totals[a] * totals[b] for a in totals for b in totals if a != b) / (n * (n - 1))
    return 1 - d_o / d_e if d_e else None


def split(claim_groups: list[str]) -> set[str]:
    """The development claims: a fixed 20%, drawn once with the protocol's seed."""
    groups = sorted(set(claim_groups))
    random.Random(SEED).shuffle(groups)  # noqa: S311 - a reproducible split, not security
    return set(groups[: round(len(groups) * DEV_SHARE)])


def evaluate(out: Path, sheets: list[Path], key_path: Path) -> dict:
    key = json.loads(key_path.read_text("utf-8"))
    examples = set(key.get("examples", {}))
    labels = read_sheets(sheets)
    overlap = {k["pair_id"] for k in key["key"] if k["overlap"]}
    alpha = krippendorff_nominal([[lab for _, lab in labels[p]] for p in overlap])
    gold = {}
    for pid, given in labels.items():
        if pid in examples:
            continue
        counts = Counter(lab for _, lab in given).most_common()
        if len(counts) > 1 and counts[0][1] == counts[1][1]:
            continue  # a tie between labellers: no gold label
        gold[pid] = counts[0][0]
    feats = {json.loads(x)["pair_id"]: json.loads(x) for x in
             (out / "matches" / "stance_features.jsonl").read_text("utf-8").splitlines()
             if x.strip()}  # fmt: skip
    ids = [p for p in gold if p in feats]
    dev_groups = split([feats[p]["claim_group"] for p in ids])
    dev = [p for p in ids if feats[p]["claim_group"] in dev_groups]
    test = [p for p in ids if feats[p]["claim_group"] not in dev_groups]
    params = fit([feats[p] for p in dev], [gold[p] for p in dev])
    g_test = [gold[p] for p in test]
    majority = Counter(gold[p] for p in dev).most_common(1)[0][0] if dev else "unrelated"
    sim_only = {**params, "entail": 2.0, "contra": 2.0, "cues": 10**6, "strong": 10**6}
    report = {
        "labelled_pairs": len(labels), "gold_pairs": len(gold), "dev": len(dev),
        "test": len(test), "alpha_on_shared_pairs": alpha, "params": params,
        "gold_distribution": dict(Counter(gold.values())),
        "test_macro_f1": {
            "model": macro_f1(g_test, [predict(feats[p], params) for p in test]),
            "majority_class": macro_f1(g_test, [majority] * len(test)),
            "similarity_only": macro_f1(g_test, [predict(feats[p], sim_only) for p in test]),
        },
    }  # fmt: skip
    (out / "matches" / "stance_eval.json").write_text(json.dumps(report, indent=1),
                                                      encoding="utf-8")  # fmt: skip
    return report
