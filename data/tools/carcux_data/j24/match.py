"""Find the news articles that may carry each claim (protocol, "Matching and stance").

Two steps, both offline:

1. `news_texts`: write the text of every collected article that is an article, has text
   and was published in the window (or has no known date), from the saved pages, to
   news/texts.jsonl. This is the corpus the matcher searches.
2. `match`: for each claim group (rows linked by same_claim_as count once), rank the
   articles published near the claim's date. The query is the claim in English together
   with the fact-checkers' own Bangla title and sentence, so Bangla articles can be found
   without translation. Ranking is BM25 on words; with `embed=True` it is fused with a
   multilingual sentence model (reciprocal rank fusion), which also finds articles that
   say the same thing in other words.

The output is a list of candidates for a person to label, never a decision: nothing here
says whether an article repeats, questions or corrects a claim. The candidate sheet shows
no verdicts, so that the stance annotator does not see them.
"""

from __future__ import annotations

import json
import math
import re
from array import array
from collections import Counter, defaultdict
from collections.abc import Callable, Iterable
from datetime import date, timedelta
from pathlib import Path

from carcux_data.j24 import WINDOW
from carcux_data.j24.news import MIN_TEXT, load_urls, read_article
from carcux_data.j24.rebuild import cached_page, offline_fetcher

BEFORE, AFTER = 7, 21  # days around a claim's date in which articles are searched
TEXT_KEEP = 5000  # characters of article text kept in texts.jsonl
SEARCHED = 3000  # characters of it searched (with the title, which counts twice)
MODEL = "intfloat/multilingual-e5-small"

_WORD = re.compile(r"[a-z0-9ঀ-৿]+")
_BN_SUFFIXES = sorted(["গুলোর", "গুলো", "দের", "েরা", "ের", "রা", "কে", "তে", "য়", "র", "ে", "ি",
                       "টি", "টা"], key=len, reverse=True)  # fmt: skip
_STOP = set("""a an the of in on at to for from by with and or but is are was were be been
has have had it its this that these those as not no into after before over said says
এবং ও এই সেই যে করে করা হয় হয়েছে হয়েছেন থেকে জন্য সঙ্গে সাথে তার তাদের এক একটি
না নয় কি কী দাবি দাবিতে শীর্ষক প্রচার প্রচারিত ছবি ভিডিও""".split())  # fmt: skip


def stem(word: str) -> str:
    """A light Bangla suffix strip (case endings and plurals); English words are kept."""
    if not re.match(r"[ঀ-৿]", word):
        return word
    for suf in _BN_SUFFIXES:
        if word.endswith(suf) and len(word) - len(suf) >= 2:
            return word[: -len(suf)]
    return word


def tokens(text: str) -> list[str]:
    words = _WORD.findall(text.lower().replace("‌", ""))
    return [stem(w) for w in words if len(w) > 1 and w not in _STOP]


def _day(value: str | None) -> date | None:
    if value and re.match(r"\d{4}-\d{2}-\d{2}", value):
        try:
            return date.fromisoformat(value[:10])
        except ValueError:
            return None
    return None


# ---------------------------------------------------------------- step 1: the corpus


def news_texts(out: Path, *, log: Callable[[str], None] = print) -> int:
    """Write news/texts.jsonl from news/articles_all.jsonl and the saved pages."""
    src = out / "news" / "articles_all.jsonl"
    urls = {u["url"]: u for u in load_urls(out)}
    fetcher = offline_fetcher(out)
    target = out / "news" / "texts.jsonl"
    tmp = target.with_suffix(".tmp")
    n = 0
    with tmp.open("w", encoding="utf-8") as f:
        for i, line in enumerate(src.read_text("utf-8").splitlines(), start=1):
            try:
                r = json.loads(line)
            except ValueError:
                continue
            if r.get("text_chars", 0) < MIN_TEXT or r.get("is_article") is False:
                continue
            if r.get("published_in_window") is False:
                continue
            body = cached_page(fetcher, urls, r)
            if not body:
                continue
            art = read_article(body, with_text=True)
            if len(art["text"]) < MIN_TEXT:
                continue
            u = urls.get(r["url"]) or {}
            capture = u.get("capture")
            f.write(json.dumps({
                "url": r["url"], "outlet": r["outlet"], "language": art.get("language"),
                "published_at": r.get("published_at"),
                "published_source": r.get("published_source"),
                "captured": f"{capture[:4]}-{capture[4:6]}-{capture[6:8]}" if capture else None,
                "title": art["title"] or r.get("title") or "",
                "text": art["text"][:TEXT_KEEP],
            }, ensure_ascii=False) + "\n")  # fmt: skip
            n += 1
            if i % 10000 == 0:
                log(f"  {i} records looked at, {n} texts written")
    tmp.replace(target)
    return n


def load_texts(out: Path) -> list[dict]:
    """The corpus, with each text cut to what is searched (to keep memory low)."""
    rows = []
    with (out / "news" / "texts.jsonl").open(encoding="utf-8") as f:
        for line in f:
            if line.strip():
                r = json.loads(line)
                r["text"] = (r.get("text") or "")[:SEARCHED]
                rows.append(r)
    return rows


# ---------------------------------------------------------------- step 2: the matcher


def claim_groups(claims_path: Path) -> list[dict]:
    """One query per claim group: the root claim, with the text of every row in it."""
    rows = [json.loads(x) for x in claims_path.read_text("utf-8").splitlines() if x.strip()]
    rows = [r for r in rows if r.get("in_scope") == "yes"]
    groups: dict[str, list[dict]] = defaultdict(list)
    for r in rows:
        groups[r["claim_group"]].append(r)
    out = []
    for gid, members in groups.items():
        root = next((m for m in members if m["claim_id"] == gid), members[0])
        days = [d for d in (_day(m.get("event_date")) for m in members) if d]
        checked = [d for d in (_day(m.get("published_at")) for m in members
                               if m["source"] == "fact_check") if d]  # fmt: skip
        ref = min(days) if days else (min(checked) - timedelta(days=BEFORE) if checked else None)
        parts = (
            [root["claim"]]
            + [m.get("title") or "" for m in members]
            + [m.get("evidence_text") or "" for m in members if m["source"] == "fact_check"]
        )
        query = re.sub(r"\[[^\]]*বাদ\]", " ", " ".join(parts))
        out.append({"claim_group": gid, "claim": root["claim"], "ref_date": ref,
                    "sources": sorted({m["fact_checker"] for m in members}),
                    "rows": len(members), "query": query})  # fmt: skip
    return out


def article_day(a: dict) -> date | None:
    return _day(a.get("published_at")) or _day(a.get("captured"))


class BM25:
    """BM25 with compact postings (doc ids and term counts in arrays), so tens of thousands
    of articles fit in a few hundred megabytes."""

    def __init__(self, docs: Iterable[list[str]], k1: float = 1.4, b: float = 0.75):
        self.k1, self.b = k1, b
        self.post: dict[str, tuple[array, array]] = {}
        lengths = array("I")
        for doc, toks in enumerate(docs):
            lengths.append(len(toks))
            for t, f in Counter(toks).items():
                ids, tfs = self.post.setdefault(t, (array("I"), array("H")))
                ids.append(doc)
                tfs.append(min(f, 65535))
        self.lengths = lengths
        self.n = len(lengths)
        self.avg = sum(lengths) / max(self.n, 1)

    def idf(self, term: str) -> float:
        f = len(self.post[term][0])
        return math.log(1 + (self.n - f + 0.5) / (f + 0.5))

    def search(self, query: list[str], allowed: set[int] | None = None) -> dict[int, float]:
        scores: dict[int, float] = defaultdict(float)
        for t in set(query):
            if t not in self.post:
                continue
            idf = self.idf(t)
            ids, tfs = self.post[t]
            for doc, f in zip(ids, tfs, strict=True):
                if allowed is not None and doc not in allowed:
                    continue
                norm = self.k1 * (1 - self.b + self.b * self.lengths[doc] / self.avg)
                scores[doc] += idf * f * (self.k1 + 1) / (f + norm)
        return scores


def _doc_tokens(a: dict) -> list[str]:
    title = tokens(a.get("title") or "")
    return title + title + tokens((a.get("text") or "")[:SEARCHED])


def _embed(texts: list[str], prefix: str, cache: Path | None, log) -> list[list[float]]:
    import numpy as np  # only with --embed
    from sentence_transformers import SentenceTransformer

    if cache and cache.exists():
        arr = np.load(cache)
        if len(arr) == len(texts):
            return arr
    log(f"  embedding {len(texts)} texts with {MODEL} (this takes a while on a CPU)")
    model = SentenceTransformer(MODEL, device="cpu")
    arr = model.encode([prefix + t for t in texts], batch_size=32, normalize_embeddings=True,
                       show_progress_bar=True, convert_to_numpy=True)  # fmt: skip
    if cache:
        np.save(cache, arr)
    return arr


def match(out: Path, claims_path: Path, *, top: int = 10, embed: bool = False,
          log: Callable[[str], None] = print) -> list[dict]:  # fmt: skip
    articles = load_texts(out)
    groups = claim_groups(claims_path)
    log(f"  {len(groups)} claim groups, {len(articles)} articles")
    index = BM25(_doc_tokens(a) for a in articles)
    days = [article_day(a) for a in articles]
    sims = None
    if embed:
        import numpy as np

        docs = _embed([f"{a['title']}. {a['text'][:1500]}" for a in articles], "passage: ",
                      out / "news" / "embed_e5small.npy", log)  # fmt: skip
        queries = _embed([g["claim"] for g in groups], "query: ", None, log)
        sims = np.asarray(queries) @ np.asarray(docs).T
    lo_w, hi_w = date.fromisoformat(WINDOW[0]), date.fromisoformat(WINDOW[1])
    results = []
    for gi, g in enumerate(groups):
        ref = g["ref_date"]
        lo = max(ref - timedelta(days=BEFORE), lo_w) if ref else lo_w
        hi = min(ref + timedelta(days=AFTER), hi_w) if ref else hi_w
        pool = [i for i, d in enumerate(days) if d is None or lo <= d <= hi]
        q = tokens(g["query"])
        found = index.search(q, set(pool))
        bm = sorted(((sc, i) for i, sc in found.items() if sc > 0), reverse=True)
        fused: dict[int, float] = defaultdict(float)
        for rank, (_, i) in enumerate(bm[:200]):
            fused[i] += 1 / (60 + rank)
        if sims is not None:
            emb = sorted(((float(sims[gi, i]), i) for i in pool), reverse=True)[:200]
            for rank, (_, i) in enumerate(emb):
                fused[i] += 1 / (60 + rank)
        bm_score = dict((i, s) for s, i in bm)
        for rank, (i, f) in enumerate(sorted(fused.items(), key=lambda kv: -kv[1])[:top], 1):
            a = articles[i]
            results.append({
                "claim_group": g["claim_group"], "rank": rank, "url": a["url"],
                "outlet": a["outlet"], "published_at": a.get("published_at"),
                "date_known": days[i] is not None, "title": a["title"],
                "bm25": round(bm_score.get(i, 0.0), 3),
                "cosine": round(float(sims[gi, i]), 4) if sims is not None else None,
                "fused": round(f, 5), "snippet": snippet(a, q),
            })  # fmt: skip
        if (gi + 1) % 50 == 0:
            log(f"  {gi + 1} of {len(groups)} claim groups matched")
    folder = out / "matches"
    folder.mkdir(exist_ok=True)
    (folder / "candidates.jsonl").write_text(
        "".join(json.dumps(r, ensure_ascii=False) + "\n" for r in results), encoding="utf-8"
    )
    write_sheet(folder / "candidates.xlsx", groups, results)
    return results


def snippet(article: dict, query: list[str], width: int = 300) -> str:
    """The 300-character stretch of the article that shares most words with the query."""
    text = " ".join((article.get("text") or "").split())
    if len(text) <= width:
        return text
    want = set(query)
    best, best_at = -1, 0
    for start in range(0, len(text) - width + 1, 100):
        hits = sum(t in want for t in tokens(text[start : start + width]))
        if hits > best:
            best, best_at = hits, start
    return text[best_at : best_at + width]


def write_sheet(path: Path, groups: list[dict], results: list[dict]) -> None:
    """The sheet a stance annotator works from: claim and article side by side, no verdict."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    claims = {g["claim_group"]: g["claim"] for g in groups}
    wb = Workbook()
    ws = wb.active
    ws.title = "Candidates"
    cols = ["pair_id", "claim_group", "claim", "rank", "outlet", "published_at", "title",
            "url", "snippet", "stance", "annotator", "note"]  # fmt: skip
    ws.append(cols)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F5F3A")
    for n, r in enumerate(results, 1):
        ws.append([f"p{n:05d}", r["claim_group"], claims[r["claim_group"]], r["rank"],
                   r["outlet"], r["published_at"], r["title"], r["url"], r["snippet"],
                   "", "", ""])  # fmt: skip
        ws.cell(ws.max_row, 8).hyperlink = r["url"]
        for c in ws[ws.max_row]:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    widths = {"claim": 45, "title": 40, "snippet": 60, "url": 18, "claim_group": 18}
    for i, name in enumerate(cols, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = widths.get(name, 12)
    dv = DataValidation(type="list", allow_blank=True,
                        formula1='"repeats,questions,corrects,mentions,unrelated"')  # fmt: skip
    dv.add(f"J2:J{ws.max_row}")
    ws.add_data_validation(dv)
    ws.freeze_panes = "D2"
    wb.save(path)
